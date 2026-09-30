"""
Tests for plugins/docs/scripts/md2docx.py — Markdown → Blue Green branded .docx.

Unit tests cover the pure helpers (front matter, placeholders, per-type options, Blue Green
page insertion). The smoke tests render the three shipped templates end to end and assert
structural invariants on the .docx (header badge, legal notice footer, cover, full-width
tables) — never exact bytes. They skip when no pandoc is reachable (PATH or pypandoc_binary).
"""
import pathlib
import subprocess
import sys
import tempfile
import unittest

from docx import Document
from docx.oxml.ns import qn

REPO_ROOT = pathlib.Path(__file__).parent.parent
PLUGIN = REPO_ROOT / "plugins/docs"
SCRIPT = PLUGIN / "scripts/md2docx.py"

sys.path.insert(0, str(SCRIPT.parent))
import md2docx  # noqa: E402

BRAND = md2docx.load_brand()
PANDOC = md2docx.find_pandoc()


class HelpersTest(unittest.TestCase):
    def test_split_front_matter(self):
        meta, body = md2docx.split_front_matter("---\ntype: report\ntitle: X\n---\n\n# A\n")
        self.assertEqual(meta, {"type": "report", "title": "X"})
        self.assertEqual(body.strip(), "# A")

    def test_no_front_matter(self):
        meta, body = md2docx.split_front_matter("# A\n")
        self.assertEqual(meta, {})
        self.assertEqual(body, "# A\n")

    def test_placeholders_read_brand_and_leave_unknown_keys(self):
        out = md2docx.fill_placeholders("{{contact.email}} {{nope.x}}", BRAND, {})
        self.assertEqual(out, f"{BRAND['contact']['email']} {{{{nope.x}}}}")

    def test_type_defaults_and_overrides(self):
        self.assertTrue(md2docx.resolve_options(BRAND, {"type": "report"})["cover"])
        self.assertFalse(md2docx.resolve_options(BRAND, {"type": "report", "cover": False})["cover"])
        self.assertTrue(md2docx.resolve_options(BRAND, {"type": "proposal"})["blue_green_page"])
        self.assertFalse(md2docx.resolve_options(BRAND, {})["toc"])

    def test_unknown_type_exits(self):
        with self.assertRaises(SystemExit):
            md2docx.resolve_options(BRAND, {"type": "memo"})

    def test_blue_green_page_at_marker_or_appended(self):
        at_marker = md2docx.insert_blue_green_page("A\n<!-- blue-green-page -->\nB", True)
        self.assertLess(at_marker.index("# Blue Green"), at_marker.index("\nB"))
        appended = md2docx.insert_blue_green_page("A", True)
        self.assertIn("# Blue Green", appended)
        self.assertNotIn("blue-green-page", md2docx.insert_blue_green_page("<!-- blue-green-page -->", False))

    def test_fields_become_content_controls(self):
        out = md2docx.expand_fields("À {{?Lieu}}, le {{?JJ/MM/AAAA#date_client}}", BRAND)
        self.assertEqual(out.count("<w:sdt>"), 2)
        self.assertIn('w:tag w:val="lieu"', out)
        self.assertIn('w:tag w:val="date_client"', out)
        self.assertIn("{=openxml}", out)
        self.assertNotIn("{{?", out)

    def test_brand_assets_exist(self):
        for name in ("badge.png", "logo.png", "logo-white.png", "cover.jpg", "blue-green-banner.jpg"):
            self.assertTrue((PLUGIN / "skills/brand/assets" / name).exists(), name)


@unittest.skipUnless(PANDOC, "pandoc not available")
class RenderSmokeTest(unittest.TestCase):
    def _render(self, template: pathlib.Path) -> Document:
        with tempfile.TemporaryDirectory() as tmp:
            src = pathlib.Path(tmp) / "doc.md"
            src.write_text(template.read_text(encoding="utf-8"), encoding="utf-8")
            out = pathlib.Path(tmp) / "doc.docx"
            result = subprocess.run([sys.executable, str(SCRIPT), str(src), "-o", str(out)],
                                    capture_output=True, text=True)
            self.assertEqual(result.returncode, 0, result.stderr)
            return Document(str(out))

    def _assert_brand_frame(self, doc):
        section = doc.sections[0]
        header_xml = section.header._element.xml
        self.assertIn("<pic:pic", header_xml, "badge missing from header")
        footer_text = "".join(p.text for p in section.footer.paragraphs)
        self.assertIn(BRAND["company"]["siret"], footer_text)
        for tbl in doc.element.body.iter(qn("w:tbl")):
            tblw = tbl.find(qn("w:tblPr")).find(qn("w:tblW"))
            self.assertEqual(tblw.get(qn("w:type")), "pct")

    def test_document(self):
        doc = self._render(PLUGIN / "skills/brand/templates/document.md")
        self._assert_brand_frame(doc)
        self.assertEqual(doc.paragraphs[0].style.name, "Title")

    def test_report_has_cover_and_toc(self):
        doc = self._render(PLUGIN / "skills/study-report/template.md")
        self._assert_brand_frame(doc)
        self.assertIn('behindDoc="1"', doc.element.body.xml, "cover image missing")
        self.assertTrue(doc.sections[0].different_first_page_header_footer)
        self.assertIn("TOC", doc.element.body.xml)

    def test_field_reaches_docx(self):
        with tempfile.TemporaryDirectory() as tmp:
            tpl = pathlib.Path(tmp) / "t.md"
            tpl.write_text("---\ntype: document\ntitle: T\n---\n\nFonction : {{?Fonction#fonction}}\n",
                           encoding="utf-8")
            doc = self._render(tpl)
        self.assertIn('w:tag w:val="fonction"', doc.element.body.xml)

    def test_proposal_has_blue_green_page(self):
        doc = self._render(PLUGIN / "skills/proposal/template.md")
        self._assert_brand_frame(doc)
        headings = [p.text for p in doc.paragraphs if p.style.name == "Heading 1"]
        self.assertIn("Blue Green", headings)
        self.assertLess(headings.index("Votre besoin"), headings.index("Blue Green"))


if __name__ == "__main__":
    unittest.main()
