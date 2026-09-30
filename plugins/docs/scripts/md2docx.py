#!/usr/bin/env python3
"""Render a Markdown file as a Blue Green branded .docx.

Usage:
    uv run --with python-docx --with pyyaml python3 scripts/md2docx.py doc.md [-o out.docx] [--pdf]

The document type (document | report | proposal) and its options come from the
Markdown front matter. Every brand value comes from skills/brand/brand.yaml.
Requires pandoc >= 2.9 on PATH, or `--with pypandoc_binary` as a fallback.
"""
import argparse
import copy
import os
import re
import shutil
import subprocess
import sys
import tempfile

try:
    import yaml
    from docx import Document
    from docx.enum.text import WD_BREAK, WD_TAB_ALIGNMENT
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls, qn
    from docx.shared import Cm, Pt, RGBColor
except ImportError as e:
    sys.exit(f"Missing dependency ({e.name}). Run with: uv run --with python-docx --with pyyaml python3 md2docx.py …")

PLUGIN = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
BRAND = os.path.join(PLUGIN, "skills", "brand")
ASSETS = os.path.join(BRAND, "assets")
PARTIAL_BG = os.path.join(BRAND, "templates", "blue-green-page.md")
MARKER_BG = "<!-- blue-green-page -->"

NEWPAGE_LUA = r"""
function RawBlock(el)
  if el.text:match('^\\newpage') or el.text:match('^\\pagebreak') then
    return pandoc.RawBlock('openxml', '<w:p><w:r><w:br w:type="page"/></w:r></w:p>')
  end
end
function HorizontalRule(el)
  return pandoc.RawBlock('openxml',
    '<w:p><w:pPr><w:pBdr><w:bottom w:val="single" w:sz="4" w:space="1" w:color="__RULE__"/></w:pBdr></w:pPr></w:p>')
end
"""


# ------------------------------------------------------------------ helpers
def load_brand():
    with open(os.path.join(BRAND, "brand.yaml"), encoding="utf-8") as f:
        return yaml.safe_load(f)


def find_pandoc():
    """pandoc on PATH first, then the binary shipped by pypandoc_binary (uv --with)."""
    exe = shutil.which("pandoc")
    if exe:
        return exe
    try:
        import pypandoc
        return pypandoc.get_pandoc_path()
    except Exception:
        return None


def split_front_matter(text):
    m = re.match(r"^---\s*\n(.*?)\n(---|\.\.\.)\s*\n", text, re.S)
    if not m:
        return {}, text
    return (yaml.safe_load(m.group(1)) or {}), text[m.end():]


def fill_placeholders(text, brand, meta):
    def repl(m):
        key = m.group(1).strip()
        node = {**brand, "doc": meta}
        for part in key.split("."):
            if not isinstance(node, dict) or part not in node:
                return m.group(0)
            node = node[part]
        return str(node)
    return re.sub(r"\{\{\s*([\w.]+)\s*\}\}", repl, text)


FIELD_RE = re.compile(r"\{\{\?\s*([^}#]+?)\s*(?:#\s*([\w-]+)\s*)?\}\}")


def slug(text):
    import unicodedata
    ascii_ = unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", "_", ascii_.lower()).strip("_") or "field"


def expand_fields(text, brand):
    """`{{?Label}}` or `{{?Label#name}}` → a Word plain-text content control showing Label.

    Word users click in the box and type; LibreOffice exports each control as a PDF form field.
    """
    c = brand["colors"]

    def repl(m):
        label, name = m.group(1), m.group(2) or slug(m.group(1))
        esc = label.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        pad = " " * max(2, (20 - len(label)) // 2)  # gives short labels a usable box width
        xml = (f'<w:sdt><w:sdtPr><w:alias w:val="{name}"/><w:tag w:val="{name}"/><w:showingPlcHdr/><w:text/>'
               f'</w:sdtPr><w:sdtContent><w:r><w:rPr><w:color w:val="{c["heading"]}"/>'
               f'<w:shd w:val="clear" w:color="auto" w:fill="{c["table_header"]}"/></w:rPr>'
               f'<w:t xml:space="preserve">{pad}{esc}{pad}</w:t></w:r></w:sdtContent></w:sdt>')
        return f"`{xml}`{{=openxml}}"
    return FIELD_RE.sub(repl, text)


def finish_pdf_form(pdf_path, brand):
    """Empty the fields LibreOffice pre-fills with their label, name them, and tint their box."""
    try:
        from pypdf import PdfReader, PdfWriter
        from pypdf.generic import (ArrayObject, BooleanObject, DictionaryObject, FloatObject,
                                   NameObject, TextStringObject)
    except ImportError:
        print("pypdf missing: PDF fields keep their label as value", file=sys.stderr)
        return 0
    reader = PdfReader(pdf_path)
    if not reader.get_fields():
        return 0
    writer = PdfWriter(clone_from=reader)
    fill, edge = brand["colors"]["table_header"], brand["colors"]["muted"]
    count = 0
    for page in writer.pages:
        for annot in page.get("/Annots") or []:
            a = annot.get_object()
            if a.get("/FT") != "/Tx":
                continue
            name = str(a.get("/TU") or a.get("/T") or "field")  # LibreOffice carries the control's alias here
            a[NameObject("/T")] = TextStringObject(name)
            a[NameObject("/V")] = TextStringObject("")
            a[NameObject("/MK")] = DictionaryObject({
                NameObject("/BG"): ArrayObject([FloatObject(int(fill[i:i + 2], 16) / 255) for i in (0, 2, 4)]),
                NameObject("/BC"): ArrayObject([FloatObject(int(edge[i:i + 2], 16) / 255) for i in (0, 2, 4)]),
            })
            if "/AP" in a:
                del a["/AP"]
            count += 1
    writer._root_object["/AcroForm"][NameObject("/NeedAppearances")] = BooleanObject(True)
    with open(pdf_path, "wb") as f:
        writer.write(f)
    return count


def rgb(hexa):
    return RGBColor.from_string(hexa.upper())


def set_run_font(rpr_parent, family):
    """Set the font on ascii/hAnsi/cs/eastAsia and drop theme font references."""
    rpr = rpr_parent.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = parse_xml(f"<w:rFonts {nsdecls('w')}/>")
        rpr.insert(0, rfonts)
    for att in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        rfonts.attrib.pop(qn(f"w:{att}"), None)
    for att in ("ascii", "hAnsi", "cs", "eastAsia"):
        rfonts.set(qn(f"w:{att}"), family)


def para_border(paragraph, side, color, sz=6, space=4):
    ppr = paragraph._p.get_or_add_pPr()
    bdr = ppr.find(qn("w:pBdr"))
    if bdr is None:
        bdr = parse_xml(f"<w:pBdr {nsdecls('w')}/>")
        ppr.append(bdr)
    bdr.append(parse_xml(
        f'<w:{side} {nsdecls("w")} w:val="single" w:sz="{sz}" w:space="{space}" w:color="{color}"/>'))


def add_field(paragraph, instr, color, size):
    run = paragraph.add_run()
    run.font.color.rgb = rgb(color)
    run.font.size = Pt(size)
    run._r.append(parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="begin"/>'))
    run2 = paragraph.add_run()
    run2._r.append(parse_xml(f'<w:instrText {nsdecls("w")} xml:space="preserve"> {instr} </w:instrText>'))
    run3 = paragraph.add_run()
    run3._r.append(parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="end"/>'))
    for r in (run2, run3):
        r.font.color.rgb = rgb(color)
        r.font.size = Pt(size)


# ------------------------------------------------------- reference.docx
def build_reference(path, brand, header_text):
    subprocess.run([PANDOC, "-o", path, "--print-default-data-file", "reference.docx"], check=True)
    doc = Document(path)
    c, pol, tz = brand["colors"], brand["font"], brand["heading_sizes_pt"]
    fam = pol["family"]

    # document default font (docDefaults) + language
    defaults = doc.styles.element.find(qn("w:docDefaults"))
    rpr_default = defaults.find(qn("w:rPrDefault")).find(qn("w:rPr"))
    for el in rpr_default.findall(qn("w:rFonts")):
        rpr_default.remove(el)
    rpr_default.insert(0, parse_xml(
        f'<w:rFonts {nsdecls("w")} w:ascii="{fam}" w:hAnsi="{fam}" w:cs="{fam}" w:eastAsia="{fam}"/>'))
    for el in rpr_default.findall(qn("w:lang")):
        rpr_default.remove(el)
    rpr_default.append(parse_xml(f'<w:lang {nsdecls("w")} w:val="fr-FR"/>'))

    # every theme font reference → brand font
    for rf in doc.styles.element.iter(qn("w:rFonts")):
        for att in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
            rf.attrib.pop(qn(f"w:{att}"), None)
        for att in ("ascii", "hAnsi", "cs", "eastAsia"):
            rf.set(qn(f"w:{att}"), fam)

    def style(name, size=None, color=None, bold=None, italic=None, before=None, after=None):
        # case-insensitive lookup: pandoc names it "Heading 1", python-docx looks for "heading 1"
        s = next((x for x in doc.styles if (x.name or "").lower() == name.lower()), None)
        if s is None:
            return None
        set_run_font(s.element, fam)
        if size:
            s.font.size = Pt(size)
        if color:
            s.font.color.rgb = rgb(color)
        if bold is not None:
            s.font.bold = bold
        if italic is not None:
            s.font.italic = italic
        pf = getattr(s, "paragraph_format", None)
        if pf is not None:
            if before is not None:
                pf.space_before = Pt(before)
            if after is not None:
                pf.space_after = Pt(after)
        return s

    for n in ("Normal", "Body Text", "First Paragraph", "Compact"):
        s = style(n, pol["body_pt"], c["text"])
        if s is not None and n != "Compact":
            s.paragraph_format.line_spacing = pol["line_spacing"]
            s.paragraph_format.space_after = Pt(6)
    h1 = style("Heading 1", tz["h1"], c["heading"], bold=False, before=24, after=10)
    style("Heading 2", tz["h2"], c["heading"], bold=False, before=16, after=6)
    style("Heading 3", tz["h3"], c["subheading"], bold=True, before=12, after=4)
    for n in ("Heading 4", "Heading 5", "Heading 6"):
        style(n, pol["body_pt"], c["subheading"], bold=True, italic=False, before=10, after=4)
    for h in (h1,):
        h.paragraph_format.keep_with_next = True
    style("Title", 24, c["heading"], bold=False, after=4)
    style("Subtitle", 13, c["subheading"], bold=False, after=4)
    style("Author", pol["body_pt"], c["muted"])
    style("Date", pol["body_pt"], c["muted"], after=18)
    style("TOC Heading", tz["h1"], c["heading"], bold=False)
    style("Image Caption", 8.5, c["muted"], italic=True)
    style("Table Caption", 8.5, c["muted"], italic=True)
    bt = style("Block Text", pol["body_pt"], c["subheading"], italic=True)
    if bt is not None:
        para_border(type("P", (), {"_p": bt.element})(), "left", c["heading"], sz=18, space=8)
    # code in a monospace font; figures centred
    for n in ("Verbatim Char", "Source Code"):
        v = next((x for x in doc.styles if (x.name or "").lower() == n.lower()), None)
        if v is not None:
            set_run_font(v.element, "Courier New")
    for n in ("Figure", "Captioned Figure", "Image Caption"):
        f = next((x for x in doc.styles if (x.name or "").lower() == n.lower()), None)
        if f is not None:
            f.paragraph_format.alignment = 1  # centred
    hl = style("Hyperlink", color=c["heading"])
    if hl is not None:
        hl.font.underline = False

    # the "Table" style pandoc applies to every table
    tbl = doc.styles.element.find(f'{qn("w:style")}[@{qn("w:styleId")}="Table"]')
    if tbl is not None:
        for tag in ("w:tblPr", "w:tblStylePr"):
            for el in tbl.findall(qn(tag)):
                tbl.remove(el)
        tbl.append(parse_xml(f"""<w:tblPr {nsdecls('w')}>
            <w:tblBorders>
              <w:top w:val="single" w:sz="4" w:color="{c['muted']}"/>
              <w:bottom w:val="single" w:sz="4" w:color="{c['muted']}"/>
              <w:insideH w:val="single" w:sz="2" w:color="{c['muted']}"/>
            </w:tblBorders>
            <w:tblCellMar><w:top w:w="60" w:type="dxa"/><w:left w:w="100" w:type="dxa"/>
              <w:bottom w:w="60" w:type="dxa"/><w:right w:w="100" w:type="dxa"/></w:tblCellMar>
          </w:tblPr>"""))
        tbl.append(parse_xml(f"""<w:tblStylePr {nsdecls('w')} w:type="firstRow">
            <w:rPr><w:b/></w:rPr>
            <w:tcPr><w:tcBorders><w:bottom w:val="single" w:sz="8" w:color="{c['heading']}"/></w:tcBorders>
              <w:shd w:val="clear" w:color="auto" w:fill="{c['table_header']}"/></w:tcPr>
          </w:tblStylePr>"""))
        tbl.append(parse_xml(f"""<w:tblStylePr {nsdecls('w')} w:type="band2Horz">
            <w:tcPr><w:shd w:val="clear" w:color="auto" w:fill="{c['table_band']}"/></w:tcPr>
          </w:tblStylePr>"""))

    # A4 page, margins, header and footer
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin = sec.bottom_margin = Cm(2.5)
    sec.left_margin = sec.right_margin = Cm(2.2)
    sec.header_distance = sec.footer_distance = Cm(1.1)
    text_width = sec.page_width - sec.left_margin - sec.right_margin

    hp = sec.header.paragraphs[0]
    hp.text = ""
    hp.paragraph_format.tab_stops.add_tab_stop(text_width, WD_TAB_ALIGNMENT.RIGHT)
    r = hp.add_run(header_text or "")
    r.font.size, r.font.color.rgb = Pt(8), rgb(c["muted"])
    hp.add_run("\t").add_picture(os.path.join(ASSETS, "badge.png"), height=Cm(0.9))
    para_border(hp, "bottom", c["logo"], sz=4, space=4)

    fp = sec.footer.paragraphs[0]
    fp.text = ""
    fp.paragraph_format.tab_stops.add_tab_stop(text_width, WD_TAB_ALIGNMENT.RIGHT)
    para_border(fp, "top", c["logo"], sz=4, space=4)
    r = fp.add_run(brand["company"]["legal_notice"] + "\t")
    r.font.size, r.font.color.rgb = Pt(7.5), rgb(c["muted"])
    add_field(fp, "PAGE", c["muted"], 8)
    fix_ppr_order(doc.styles.element)
    fix_ppr_order(sec.header._element)
    fix_ppr_order(sec.footer._element)
    fix_ppr_order(doc.element)
    doc.save(path)


# ------------------------------------------------------- post-processing
def to_anchor_behind(run, width_emu, height_emu):
    """Turn a run's inline image into an image anchored to the page, behind the text."""
    inline = run._r.find(".//" + qn("wp:inline"))
    graphic = inline.find(qn("a:graphic"))
    docpr = inline.find(qn("wp:docPr"))
    anchor = parse_xml(
        f'<wp:anchor {nsdecls("wp", "a", "pic", "r")} distT="0" distB="0" distL="0" distR="0" '
        f'simplePos="0" relativeHeight="0" behindDoc="1" locked="1" layoutInCell="1" allowOverlap="1">'
        f'<wp:simplePos x="0" y="0"/>'
        f'<wp:positionH relativeFrom="page"><wp:posOffset>0</wp:posOffset></wp:positionH>'
        f'<wp:positionV relativeFrom="page"><wp:posOffset>0</wp:posOffset></wp:positionV>'
        f'<wp:extent cx="{width_emu}" cy="{height_emu}"/><wp:effectExtent l="0" t="0" r="0" b="0"/>'
        f'<wp:wrapNone/></wp:anchor>')
    anchor.append(copy.deepcopy(docpr))
    anchor.append(parse_xml(f'<wp:cNvGraphicFramePr {nsdecls("wp")}/>'))
    anchor.append(copy.deepcopy(graphic))
    inline.getparent().replace(inline, anchor)


def add_cover(doc, brand, meta):
    c = brand["colors"]
    sec = doc.sections[0]
    sec.different_first_page_header_footer = True
    sec.first_page_header.is_linked_to_previous = False
    sec.first_page_footer.is_linked_to_previous = False

    body = doc.element.body
    first = body[0]
    paras = []

    def new_par(before=0):
        p = doc.add_paragraph()  # appended, then moved to the top
        p.paragraph_format.space_before = Pt(before)
        p.paragraph_format.space_after = Pt(0)
        p.style = doc.styles["Normal"]
        paras.append(p)
        return p

    def white(p, text, size, bold=False):
        r = p.add_run(text)
        r.font.size, r.font.bold = Pt(size), bold
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        return r

    p = new_par()
    r = p.add_run()
    r.add_picture(os.path.join(ASSETS, "cover.jpg"), width=sec.page_width, height=sec.page_height)
    to_anchor_behind(r, sec.page_width, sec.page_height)

    p = new_par(before=190)
    p.add_run().add_picture(os.path.join(ASSETS, "logo-white.png"), width=Cm(5.5))
    p = new_par(before=6)
    para_border(p, "bottom", "FFFFFF", sz=6, space=1)
    p = new_par(before=18)
    white(p, meta.get("title", ""), 24)
    if meta.get("subtitle"):
        p = new_par(before=6)
        white(p, meta["subtitle"], 13)

    p = new_par(before=200)
    para_border(p, "top", "FFFFFF", sz=6, space=6)
    text_width = sec.page_width - sec.left_margin - sec.right_margin
    p.paragraph_format.tab_stops.add_tab_stop(text_width, WD_TAB_ALIGNMENT.RIGHT)
    white(p, str(meta.get("date", "")), 9)
    white(p, "\t" + str(meta.get("client", "")), 9)
    paras[-1].add_run().add_break(WD_BREAK.PAGE)

    for p in paras:
        first.addprevious(p._p)


def page_break_after_toc(doc):
    body = doc.element.body
    for sdt in body.iter(qn("w:sdt")):
        if sdt.find(".//" + qn("w:docPartGallery")) is not None or "TOC" in "".join(
                t.text or "" for t in sdt.iter(qn("w:instrText"))):
            sdt.addnext(parse_xml(f'<w:p {nsdecls("w")}><w:r><w:br w:type="page"/></w:r></w:p>'))
            return


PPR_ORDER = ["pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl", "numPr",
             "suppressLineNumbers", "pBdr", "shd", "tabs", "suppressAutoHyphens", "kinsoku", "wordWrap",
             "overflowPunct", "topLinePunct", "autoSpaceDE", "autoSpaceDN", "bidi", "adjustRightInd",
             "snapToGrid", "spacing", "ind", "contextualSpacing", "mirrorIndents", "suppressOverlap", "jc",
             "textDirection", "textAlignment", "textboxTightWrap", "outlineLvl", "divId", "cnfStyle",
             "rPr", "sectPr", "pPrChange"]


def fix_ppr_order(root):
    """Reorder every w:pPr's children to the OOXML schema order (Word is strict)."""
    rank = {qn(f"w:{n}"): i for i, n in enumerate(PPR_ORDER)}
    for ppr in root.iter(qn("w:pPr")):
        kids = list(ppr)
        kids.sort(key=lambda el: rank.get(el.tag, len(rank)))
        for el in kids:
            ppr.append(el)
    for mar in root.iter(qn("w:pgMar")):
        mar.set(qn("w:gutter"), mar.get(qn("w:gutter"), "0"))


def normalize_tables(doc):
    """Full-width tables that keep their column proportions (old pandoc versions break widths)."""
    sec = doc.sections[0]
    width_twips = int((sec.page_width - sec.left_margin - sec.right_margin) / 635)
    for tbl in doc.element.body.iter(qn("w:tbl")):
        tblpr = tbl.find(qn("w:tblPr"))
        for el in tblpr.findall(qn("w:tblW")) + tblpr.findall(qn("w:tblLayout")):
            tblpr.remove(el)
        style_el = tblpr.find(qn("w:tblStyle"))
        pos = 1 if style_el is not None else 0
        tblpr.insert(pos, parse_xml(f'<w:tblW {nsdecls("w")} w:w="5000" w:type="pct"/>'))
        grid = tbl.find(qn("w:tblGrid"))
        cols = grid.findall(qn("w:gridCol")) if grid is not None else []
        if cols:
            # keep the proportions pandoc derived from the Markdown dashes, scaled to full width
            ws = [int(gc.get(qn("w:w")) or 0) for gc in cols]
            if not all(ws) or len(set(ws)) == 1:
                ws = [1] * len(cols)
            tot = sum(ws)
            for gc, w in zip(cols, ws):
                gc.set(qn("w:w"), str(width_twips * w // tot))
        # keep short tables on one page (keepNext on every row but the last)
        rows = tbl.findall(qn("w:tr"))
        if len(rows) <= 15:
            for tr in rows[:-1]:
                for p in tr.iter(qn("w:p")):
                    ppr = p.find(qn("w:pPr"))
                    if ppr is None:
                        ppr = parse_xml(f'<w:pPr {nsdecls("w")}/>')
                        p.insert(0, ppr)
                    if ppr.find(qn("w:keepNext")) is None:
                        ppr.append(parse_xml(f'<w:keepNext {nsdecls("w")}/>'))
        for tcw in list(tbl.iter(qn("w:tcW"))):
            tcw.set(qn("w:type"), "auto")
            tcw.set(qn("w:w"), "0")


def fix_all_parts(doc):
    normalize_tables(doc)
    fix_ppr_order(doc.element)
    fix_ppr_order(doc.styles.element)
    for sec in doc.sections:
        for hf in (sec.header, sec.footer, sec.first_page_header, sec.first_page_footer):
            if not hf.is_linked_to_previous:
                fix_ppr_order(hf._element)


# ------------------------------------------------------------------ main
PANDOC = None


def resolve_options(brand, meta):
    kind = meta.get("type", "document")
    if kind not in brand["types"]:
        sys.exit(f"unknown type: {kind} (expected: {', '.join(brand['types'])})")
    defaults = brand["types"][kind]
    return {**defaults, **{k: meta[k] for k in defaults if k in meta}}


def insert_blue_green_page(body, enabled):
    if not enabled:
        return body.replace(MARKER_BG, "")
    with open(PARTIAL_BG, encoding="utf-8") as f:
        partial = f.read()
    return body.replace(MARKER_BG, partial) if MARKER_BG in body else body.rstrip() + "\n\n" + partial


def main():
    global PANDOC
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("source")
    ap.add_argument("-o", "--output")
    ap.add_argument("--pdf", action="store_true", help="also write a PDF (LibreOffice) for visual checks")
    args = ap.parse_args()

    PANDOC = find_pandoc()
    if not PANDOC:
        sys.exit("pandoc not found. Install it, or add `--with pypandoc_binary` to the uv command.")

    brand = load_brand()
    src = os.path.abspath(args.source)
    out = os.path.abspath(args.output or os.path.splitext(src)[0] + ".docx")
    with open(src, encoding="utf-8") as f:
        meta, body = split_front_matter(f.read())
    opts = resolve_options(brand, meta)
    body = insert_blue_green_page(body, opts["blue_green_page"])
    body = expand_fields(fill_placeholders(body, brand, meta), brand)

    with tempfile.TemporaryDirectory() as tmp:
        ref = os.path.join(tmp, "reference.docx")
        build_reference(ref, brand, meta.get("header") or meta.get("title", ""))
        md = os.path.join(tmp, "source.md")
        with open(md, "w", encoding="utf-8") as f:
            f.write(body)
        lua = os.path.join(tmp, "newpage.lua")
        with open(lua, "w") as f:
            f.write(NEWPAGE_LUA.replace("__RULE__", brand["colors"]["muted"]))

        cmd = [PANDOC, md, "-f", "markdown+raw_tex", "-o", out, "--reference-doc", ref,
               "--lua-filter", lua,
               "--resource-path", os.pathsep.join([os.path.dirname(src), BRAND]),
               "-M", "lang=fr-FR"]
        if not opts["cover"]:  # pandoc title block at the top of page 1
            for k in ("title", "subtitle", "date"):
                if meta.get(k):
                    cmd += ["-M", f"{k}={meta[k]}"]
        if opts["toc"]:
            cmd += ["--toc", "--toc-depth=2", "-M", "toc-title=Sommaire"]
        subprocess.run(cmd, check=True)

    doc = Document(out)
    if opts["toc"]:
        page_break_after_toc(doc)
    if opts["cover"]:
        add_cover(doc, brand, meta)
    fix_all_parts(doc)
    doc.save(out)
    print(out)

    if args.pdf:
        soffice = shutil.which("soffice") or shutil.which("libreoffice")
        if not soffice:
            print("soffice not found: PDF skipped", file=sys.stderr)
            return
        # ExportFormFields turns every Word content control ({{?…}}) into a fillable PDF field
        pdf_filter = 'pdf:writer_pdf_Export:{"ExportFormFields":{"type":"boolean","value":"true"}}'
        subprocess.run([soffice, "--headless", "--convert-to", pdf_filter, "--outdir",
                        os.path.dirname(out), out], check=True, capture_output=True)
        pdf = os.path.splitext(out)[0] + ".pdf"
        n = finish_pdf_form(pdf, brand)
        print(pdf + (f"  ({n} fillable fields)" if n else ""))


if __name__ == "__main__":
    main()
